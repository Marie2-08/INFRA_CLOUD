"""Opdracht 3, 4 en 5: de gefilterde dict opslaan als JSON, CSV en XLSX (en weer inlezen)."""
import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCHEIDINGSTEKEN = ";"     # Nederlandstalige/Belgische Excel verwacht ; (komma = decimaalteken)
DECIMAAL = ","
EXTRA_KOLOM = "bron_bestand"
BLADNAAM = "Systeeminfo"
LETTERTYPE = "Arial"

# Tabelstructuur: (kolomnaam, pad in de dict, type)
KOLOMMEN = [
    ("computernaam",            ("device_name",), str),
    ("computermodel",           ("device_model",), str),
    ("serienummer",             ("serial_number",), str),
    ("processor",               ("processor",), str),
    ("fysieke_cpu_cores",       ("physical_cpu_cores",), int),
    ("max_cpu_frequentie_mhz",  ("maximum_cpu_frequency_mhz",), float),
    ("totaal_ram_gib",          ("total_ram_gib",), float),
    ("bestandssysteem",         ("storage", "file_system"), str),
    ("opslagcapaciteit_gib",    ("storage", "capacity_gib"), float),
    ("netwerkinterface",        ("primary_network_interface", "interface"), str),
    ("mac_adres",               ("primary_network_interface", "mac_address"), str),
    ("ipv4_adres",              ("primary_network_interface", "ipv4_address"), str),
    ("ipv6_adres",              ("primary_network_interface", "ipv6_address"), str),
    ("ipv6_link_local_adres",   ("primary_network_interface", "ipv6_link_local_address"), str),
]
KOLOMNAMEN = [naam for naam, _, _ in KOLOMMEN]
TYPES = {naam: typ for naam, _, typ in KOLOMMEN}


def veilige_naam(tekst):
    """Maak van een computernaam een veilige bestandsnaam."""
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", (tekst or "").strip()) or "onbekend"


# ---------------------------------------------------------------------------
# Van geneste dict naar een rij (één computer = één rij)
# ---------------------------------------------------------------------------
def haal_op(info, pad):
    """Haal een (geneste) waarde uit de dict; None als ze ontbreekt."""
    waarde = info
    for sleutel in pad:
        if not isinstance(waarde, dict):
            return None
        waarde = waarde.get(sleutel)
    return waarde


def naar_tekst(waarde):
    """Ontbrekende waarde -> leeg veld; getal -> tekst met het gekozen decimaalteken."""
    if waarde is None:
        return ""
    if isinstance(waarde, float):
        return str(waarde).replace(".", DECIMAAL)
    return str(waarde)


def dict_naar_rij(info):
    return {naam: naar_tekst(haal_op(info, pad)) for naam, pad, _ in KOLOMMEN}


def tekst_naar_waarde(tekst, typ):
    """Tekst uit een CSV-rij -> Python-waarde met het juiste type (leeg = None)."""
    if tekst is None or tekst == "":
        return None
    if typ is float:
        return float(str(tekst).replace(",", "."))
    if typ is int:
        return int(round(float(str(tekst).replace(",", "."))))
    return str(tekst)


# ---------------------------------------------------------------------------
# JSON (opdracht 3)
# ---------------------------------------------------------------------------
def schrijf_json(info, bestand):
    with open(bestand, "w", encoding="utf-8") as f:
        json.dump(info, f, indent=2, ensure_ascii=False)


def lees_json(bestand):
    with open(bestand, "r", encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# CSV (opdracht 4)
# ---------------------------------------------------------------------------
def schrijf_csv(rijen, bestand):
    """Schrijf rijen (lijst van tekst-dicts) naar een CSV-bestand dat Excel goed opent."""
    kolommen = list(KOLOMNAMEN)
    if rijen and EXTRA_KOLOM in rijen[0]:
        kolommen.append(EXTRA_KOLOM)
    # utf-8-sig: Excel herkent de tekens correct; newline="": geen lege regels op Windows
    with open(bestand, "w", newline="", encoding="utf-8-sig") as f:
        schrijver = csv.DictWriter(f, fieldnames=kolommen, delimiter=SCHEIDINGSTEKEN)
        schrijver.writeheader()
        schrijver.writerows(rijen)


def lees_csv(bestand):
    """Lees een CSV-bestand terug als lijst van dicts met de juiste types."""
    rijen = []
    with open(bestand, newline="", encoding="utf-8-sig") as f:
        for rij in csv.DictReader(f, delimiter=SCHEIDINGSTEKEN):
            rijen.append({naam: tekst_naar_waarde(tekst, TYPES.get(naam, str))
                          for naam, tekst in rij.items()})
    return rijen


# ---------------------------------------------------------------------------
# XLSX (opdracht 5)
# ---------------------------------------------------------------------------
def schrijf_xlsx(rijen, bestand):
    """Schrijf rijen (lijst van tekst-dicts) naar een opgemaakt Excel-bestand met echte getallen."""
    kolommen = list(KOLOMNAMEN)
    if rijen and EXTRA_KOLOM in rijen[0]:
        kolommen.append(EXTRA_KOLOM)

    wb = Workbook()
    ws = wb.active
    ws.title = BLADNAAM
    ws.append(kolommen)
    for rij in rijen:
        ws.append([tekst_naar_waarde(rij.get(naam, ""), TYPES.get(naam, str)) for naam in kolommen])

    for cel in ws[1]:
        cel.font = Font(name=LETTERTYPE, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", start_color="305496")
        cel.alignment = Alignment(horizontal="center", vertical="center")

    for rij_cellen in ws.iter_rows(min_row=2):
        for cel, naam in zip(rij_cellen, kolommen):
            cel.font = Font(name=LETTERTYPE)
            typ = TYPES.get(naam, str)
            if typ is int:
                cel.number_format = "0"
            elif typ is float:
                cel.number_format = "0.00"
            elif naam == "serienummer":
                cel.number_format = "@"   # tekstformaat: Excel mag er geen getal van maken

    for kolom in ws.columns:
        breedte = max(len(str(c.value)) if c.value is not None else 0 for c in kolom)
        ws.column_dimensions[kolom[0].column_letter].width = breedte + 3

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(bestand)


def lees_xlsx(bestand):
    ws = load_workbook(bestand)[BLADNAAM]
    koppen = [c.value for c in ws[1]]
    return [dict(zip(koppen, rij)) for rij in ws.iter_rows(min_row=2, values_only=True)]
